import docx

doc = docx.Document("synopsis.docx")
print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs[:50]):
    if p.text.strip():
        print(f"P{i} [Style={p.style.name}]: {p.text[:100]}...")
