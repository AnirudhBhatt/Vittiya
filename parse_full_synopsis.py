import docx

doc = docx.Document("synopsis.docx")
with open("synopsis_dump.txt", "w", encoding="utf-8") as f:
    f.write(f"Total Paragraphs: {len(doc.paragraphs)}\n")
    f.write(f"Total Tables: {len(doc.tables)}\n\n")
    for i, p in enumerate(doc.paragraphs):
        f.write(f"[{i}] [Style={p.style.name}] (sb={p.paragraph_format.space_before}, sa={p.paragraph_format.space_after}, ls={p.paragraph_format.line_spacing}): {p.text}\n")

print("Dumped synopsis.docx to synopsis_dump.txt")
