from pptx import Presentation

prs = Presentation("Vittiya PPT.pptx")
print(f"Total Slides: {len(prs.slides)}")

for idx, slide in enumerate(prs.slides, start=1):
    print(f"\n==========================================")
    print(f"SLIDE {idx}")
    print(f"==========================================")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if p.text.strip():
                    print(f"  • {p.text.strip()}")
