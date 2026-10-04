import docx

doc = docx.Document("Major (1).docx")
print(f"Total Paragraphs in Major (1).docx: {len(doc.paragraphs)}")
print(f"Total Tables in Major (1).docx: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs):
    if any(k in p.text for k in ["Methodology", "PERT", "Gantt", "Timeline", "Phase", "2.1", "2.2", "2.3", "1.3", "Schedule", "WBS"]):
        print(f"P[{i}]: {p.text[:120]}")
